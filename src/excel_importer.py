"""Read the documented Excel template and reuse the configuration compiler."""
import argparse
import json
from pathlib import Path
from zipfile import BadZipFile

try:
    from .generator import generate, require, ValidationError
except ImportError:
    from generator import generate, require, ValidationError

HEADERS = {
    'Model': ['model_id', 'model_name'],
    'Factors': ['id', 'label', 'type', 'unit', 'weight_percent'],
    'Rules': ['factor_id', 'min', 'max', 'value', 'score'],
}
MAX_ROWS = 1005


def build_input(tables):
    """Map table rows to the same input contract used by JSON callers."""
    model = tables['Model']
    require(len(model) == 1, 'Model requires one data row at row 6')
    source = dict(zip(HEADERS['Model'], model[0]))
    factors, by_id = [], {}
    for row_number, row in tables['Factors']:
        factor = dict(zip(HEADERS['Factors'], row))
        identifier = factor['id']
        require(isinstance(identifier, str) and identifier.strip(), f'Factors!A{row_number}: missing text ID')
        require(identifier not in by_id, f'Factors!A{row_number}: duplicate ID {identifier}')
        factor['rules'] = []
        factors.append(factor)
        by_id[identifier] = factor
    for row_number, row in tables['Rules']:
        identifier, lower, upper, value, score = row
        require(isinstance(identifier, str) and identifier in by_id,
                f'Rules!A{row_number}: unknown factor ID {identifier!r}')
        factor = by_id[identifier]
        if factor['type'] == 'numeric':
            require(value is None, f'Rules!D{row_number}: numeric rules must leave value blank')
            rule = {'min': lower, 'max': upper, 'score': score}
        else:
            require(lower is None and upper is None, f'Rules!B{row_number}: categorical rules must leave min/max blank')
            rule = {'value': value, 'score': score}
        factor['rules'].append(rule)
    source['factors'] = factors
    return source


def read_tables(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError('Excel import requires: python -m pip install -r requirements.txt') from exc
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except (OSError, ValueError, BadZipFile, KeyError) as exc:
        raise ValidationError(f'Cannot read Excel workbook: {exc}') from exc
    try:
        require(set(workbook.sheetnames) == set(HEADERS), 'Workbook must contain exactly Model, Factors, and Rules sheets')
        tables = {}
        for name, headers in HEADERS.items():
            sheet = workbook[name]
            require(sheet.max_column <= len(headers), f'{name}: unexpected columns outside the input table')
            actual = [cell.value for cell in sheet[5]][:len(headers)]
            require(actual == headers, f'{name}: row 5 headers must be {headers}')
            require(sheet.max_row <= MAX_ROWS, f'{name}: input supports rows 6 through {MAX_ROWS}')
            last = 6 if name == 'Model' else sheet.max_row
            rows = []
            for cells in sheet.iter_rows(min_row=6, max_row=last, max_col=len(headers)):
                values = [cell.value for cell in cells]
                if all(value is None for value in values):
                    continue
                for cell in cells:
                    require(cell.data_type not in ('f', 'e'), f'{name}!{cell.coordinate}: use literal inputs, not formulas or Excel errors')
                rows.append(values if name == 'Model' else (cells[0].row, values))
            tables[name] = rows
        return tables
    finally:
        workbook.close()


def load_input(path):
    return build_input(read_tables(path))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try:
        require(args.input.suffix.lower() == '.xlsx', 'Excel input must be an .xlsx file')
        result = generate(load_input(args.input))
        payload = json.dumps(result, indent=2, allow_nan=False) + '\n'
        args.output.write_text(payload, encoding='utf-8')
    except (ValidationError, ValueError, OSError) as exc:
        parser.exit(1, f'Error: {exc}\n')
    print(f'Generated {args.output} from Excel with {len(result["factors"])} validated factors')


if __name__ == '__main__':
    main()
